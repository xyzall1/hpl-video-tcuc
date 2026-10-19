"""Put rendered thumbnails into PowerPoint, one full-bleed image per slide, notes = file name.

  python3 make_pptx.py <out_prefix> <png_dir> [<png_dir> ...]
    -> <out_prefix>_Thumbnail_16x9.pptx        video, artikel, cover course, cover kompetensi
    -> <out_prefix>_Thumbnail_1x1_Audio.pptx   audio

Give one folder per level (L1, L2, L3) to get ONE deck for all levels. Within each level the 16:9 slides follow
the CMS column order: Judul Video (all LUs) -> Judul Artikel (all LUs) -> Nama Course (course covers)
-> Learning Path (competency cover). Audio slides follow the LU order."""
import sys, os, re, glob
from pptx import Presentation
from pptx.util import Emu

def lu(b):
    m = re.search(r"_KB(\d+)_M(\d+)_", b)
    return (int(m.group(1)), int(m.group(2))) if m else (0, 0)

def level(b):
    m = re.search(r"_L(\d+)_", b)
    return int(m.group(1)) if m else 0

def wide_key(f):
    b = os.path.basename(f)
    if "_Video - " in b: grp = 0
    elif "_Artikel - " in b: grp = 1
    elif "_Cover Course " in b: grp = 2
    else: grp = 3                                   # cover kompetensi = learning path
    cn = re.search(r"_Cover Course (\d+)", b)
    return (level(b), grp, lu(b), int(cn.group(1)) if cn else 0, b)

def deck(files, W, H, path):
    pr = Presentation(); pr.slide_width, pr.slide_height = Emu(W), Emu(H)
    for f in files:
        sl = pr.slides.add_slide(pr.slide_layouts[6]); sl.shapes.add_picture(f, 0, 0, Emu(W), Emu(H))
        sl.notes_slide.notes_text_frame.text = os.path.basename(f)[:-4]
    pr.save(path); print("pptx:", os.path.abspath(path), len(files), "slide")

def main(prefix, *dirs):
    pngs = [f for d in dirs for f in glob.glob(os.path.join(d, "*.png"))]
    wide = sorted([f for f in pngs if "_Audio - " not in os.path.basename(f)], key=wide_key)
    audio = sorted([f for f in pngs if "_Audio - " in os.path.basename(f)],
                   key=lambda f: (level(os.path.basename(f)), lu(os.path.basename(f)), f))
    if wide: deck(wide, 12192000, 6858000, f"{prefix}_Thumbnail_16x9.pptx")
    if audio: deck(audio, 6858000, 6858000, f"{prefix}_Thumbnail_1x1_Audio.pptx")

if __name__ == "__main__":
    main(*sys.argv[1:])
