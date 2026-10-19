"""Build the HPL CMS upload workbook (one sheet per level, one row per LU).

  python3 build_cms.py build <data.json> <out.xlsx>

data.json:
{
  "competency": "User Interface Design",
  "credits": {"author": "...", "reviewer": "...", "video_tool": "Claude", "audio_tool": "NotebookLM"},   # always ask the user
  "sheets": [
    {"name": "TC229 L1", "level": "Fundamental",
     "courses": {"KB1": {"name": "<course plan name, '' if unknown>",
                         "kb": "prinsip-prinsip desain user interface."}},   # KB text after the opening verb
     "lus": [{"kb": "KB1", "modul": 1, "judul_modul": "...", "judul_video": "...", "caption_video": "...",
              "judul_audio": "...", "caption_audio": "...", "judul_artikel": "..."}]}
  ]
}
Captions are the content sentences only; the credit block is appended after them (description first, then a blank line, then credits). Empty required cells are highlighted yellow."""
import sys, json, re, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

COLS = [("Learning Content Title", 30), ("Text Core Learning Content", 50), ("Judul Video", 32), ("Deskripsi Video", 60),
        ("Judul Audio", 32), ("Deskripsi Audio", 60), ("Judul Artikel", 36), ("Section Title", 30),
        ("Nama Course", 30), ("Deskripsi Course", 50), ("Learning Path Title", 28), ("Learning Path Description", 40)]
TEXT_CORE = "Konten ini berfokus pada pembahasan mengenai "
COURSE_DESC = "Course ini berfokus pada pembahasan mengenai "

def credit(tool, c):
    return (f"Produced with the support of Artificial Intelligence (AI) tools: {tool}.\n"
            f"Learning content created by {c['author']}.\nReviewed by {c['reviewer']}.")

def check(d):
    warn = []
    c = d.get("credits") or {}
    for k in ("author", "reviewer", "video_tool", "audio_tool"):
        if not c.get(k): warn.append(f"credits.{k} kosong: tanyakan ke user")
    for sh in d["sheets"]:
        for kb, co in sh["courses"].items():
            if not co.get("name"): warn.append(f"{sh['name']} {kb}: Nama Course belum ada (kolom dikosongkan, ditandai kuning)")
            if re.match(r"(?i)(mampu|menerapkan|mengidentifikasi|menganalisis|merancang|melakukan|menjalankan)\b", co["kb"]):
                warn.append(f"{sh['name']} {kb}: teks KB masih diawali kata kerja; ubah agar nyambung setelah 'pembahasan mengenai'")
        for lu in sh["lus"]:
            tag = f"{sh['name']} {lu['kb']} M{lu['modul']}"
            for k in ("caption_video", "caption_audio"):
                t = lu.get(k, "")
                if t.startswith("Produced with"): warn.append(f"{tag}: {k} sudah berisi kredit; tulis isi saja")
                if not re.search(r"\b(Yuk|Mari)\b", t): warn.append(f"{tag}: {k} belum bernada ajakan (Yuk/Mari)")
            if re.search(r"\[[^\]]+\]", " ".join(str(v) for v in lu.values())): warn.append(f"{tag}: ada tanda [..] tersisa")
    return warn

def build(data, out):
    d = json.load(open(data)); c = d["credits"]
    for w in check(d): print("CEK:", w)
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    thin = Side(style="thin", color="FFBFBFBF"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
    yellow = PatternFill("solid", fgColor="FFFFF2CC")
    for sh in d["sheets"]:
        ws = wb.create_sheet(sh["name"][:31])
        for i, (h, wdt) in enumerate(COLS, 1):
            cell = ws.cell(1, i, h); cell.border = box
            cell.font = Font(bold=True, color="FFFFFFFF"); cell.fill = PatternFill("solid", fgColor="FF2796A3")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = wdt
        ws.row_dimensions[1].height = 32
        lus = sorted(sh["lus"], key=lambda u: (int(re.sub(r"\D", "", u["kb"])), int(u["modul"])))
        for r, lu in enumerate(lus, 2):
            co = sh["courses"][lu["kb"]]
            row = [lu["judul_modul"], TEXT_CORE + co["kb"], lu["judul_video"],
                   lu["caption_video"].strip() + "\n\n" + credit(c["video_tool"], c), lu["judul_audio"],
                   lu["caption_audio"].strip() + "\n\n" + credit(c["audio_tool"], c), lu["judul_artikel"],
                   lu["judul_modul"], co.get("name", ""), COURSE_DESC + co["kb"],
                   f"{d['competency']} {sh['level']}", f"Learning Path for {d['competency']}: {sh['level']}"]
            for i, v in enumerate(row, 1):
                cell = ws.cell(r, i, v); cell.border = box
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if not v: cell.fill = yellow
            ws.row_dimensions[r].height = 150
        ws.freeze_panes = "B2"
    wb.save(out)
    print("tersimpan:", out, "|", ", ".join(f"{s['name']} ({len(s['lus'])} LU)" for s in d["sheets"]))

if __name__ == "__main__":
    {"build": build}[sys.argv[1]](*sys.argv[2:4])
